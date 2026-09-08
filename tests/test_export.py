from pathlib import Path
from tempfile import TemporaryDirectory

from openpyxl import Workbook, load_workbook

from overnight_scraper.export import export_master_copy, export_master_csv
from overnight_scraper.models import ColorMedia, Family, FamilyResult


def _make_master(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Overnight Products"
    sheet.append([
        "SKU",
        "Default_image_url",
        "Images Link",
        "Videos Link",
        "Yellow Gold Images Link",
        "Yellow Gold Videos Link",
        "Rose Gold Images Link",
        "Rose Gold Videos Link",
    ])
    sheet.append(["SKU-1", "old-default", "old-white", "old-video", "old-yellow", "old-yvideo", "old-rose", "old-rvideo"])
    workbook.save(path)
    workbook.close()


def test_exports_write_new_files_and_preserve_source():
    with TemporaryDirectory() as directory:
        root = Path(directory)
        source = root / "master.xlsx"
        xlsx_out = root / "downloads" / "overnight_latest.xlsx"
        csv_out = root / "downloads" / "overnight_latest.csv"
        _make_master(source)
        families = [Family("FAM-1", "SKU-1", ("SKU-1",), 1)]
        results = [
            FamilyResult(
                family_id="FAM-1",
                source_sku="SKU-1",
                product_url="https://example.test/product/sku-1",
                white=ColorMedia(["new-white.jpg"], [], True, {"color_specific": True}),
                yellow_gold=ColorMedia(["new-yellow.jpg"], [], True, {"color_specific": True}),
                rose_gold=ColorMedia([], [], True, {"color_specific": False}),
            )
        ]

        export_master_copy(source, xlsx_out, families, results)
        export_master_csv(source, csv_out, families, results)

        source_workbook = load_workbook(source)
        try:
            assert source_workbook["Overnight Products"]["B2"].value == "old-default"
            assert source_workbook["Overnight Products"]["G2"].value == "old-rose"
        finally:
            source_workbook.close()

        output_workbook = load_workbook(xlsx_out)
        try:
            sheet = output_workbook["Overnight Products"]
            assert sheet["B2"].value == "new-white.jpg"
            assert "new-yellow.jpg" in sheet["E2"].value
            assert sheet["G2"].value == "old-rose"
        finally:
            output_workbook.close()

        csv_text = csv_out.read_text(encoding="utf-8-sig")
        assert "new-white.jpg" in csv_text
        assert "new-yellow.jpg" in csv_text
        assert "old-rose" in csv_text
