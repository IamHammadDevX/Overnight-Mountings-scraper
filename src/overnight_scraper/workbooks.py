from __future__ import annotations

from collections import defaultdict
from pathlib import Path
from openpyxl import load_workbook
from .models import Family

FAMILIES_SHEET = "Families To Scrape"
UNIQUE_SKUS_SHEET = "All Unique SKUs"
MASTER_SHEET = "Overnight Products"

def sku_text(value: object) -> str:
    return "" if value is None else (value if isinstance(value, str) else str(value))

def _headers(sheet) -> dict[str, int]:
    return {sku_text(cell.value): index for index, cell in enumerate(next(sheet.iter_rows()), 1)}

def load_families(requirements_path: Path) -> list[Family]:
    """Read requirements workbook only and validate its dynamic family counts."""
    workbook = load_workbook(requirements_path, read_only=True, data_only=False)
    try:
        worklist, unique = workbook[FAMILIES_SHEET], workbook[UNIQUE_SKUS_SHEET]
        wh, uh = _headers(worklist), _headers(unique)
        needed = {"Family ID", "# SKUs in Family", "Representative SKU"}
        if missing := needed - wh.keys():
            raise ValueError(f"Missing worklist columns: {sorted(missing)}")
        members: dict[str, list[str]] = defaultdict(list)
        for row in unique.iter_rows(min_row=2, values_only=True):
            sku, family_id = sku_text(row[uh["SKU"] - 1]), sku_text(row[uh["Family ID"] - 1])
            if sku and family_id:
                members[family_id].append(sku)
        output = []
        for row in worklist.iter_rows(min_row=2, values_only=True):
            family_id = sku_text(row[wh["Family ID"] - 1])
            if not family_id:
                continue
            representative = sku_text(row[wh["Representative SKU"] - 1])
            declared = row[wh["# SKUs in Family"] - 1]
            skus = members.get(family_id, [])
            if not isinstance(declared, int) or len(skus) != declared or representative not in skus:
                raise ValueError(f"Invalid family mapping: {family_id}")
            # Required fallback sequence: representative, then remaining source-row order.
            output.append(Family(family_id, representative, tuple([representative, *(sku for sku in skus if sku != representative)]), declared))
        return output
    finally:
        workbook.close()
